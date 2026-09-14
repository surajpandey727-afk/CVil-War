"""Resume and document parser supporting PDF and DOCX formats."""

import asyncio
import re
from pathlib import Path

import structlog
from pydantic import BaseModel, ConfigDict, Field

from app.core.exceptions import ParseError

logger = structlog.get_logger(__name__)

_SECTION_HEADERS: list[str] = [
    "summary", "objective", "professional summary", "profile",
    "experience", "work experience", "professional experience", "employment",
    "education", "academic background",
    "skills", "technical skills", "core competencies", "competencies",
    "certifications", "certificates", "licenses",
    "projects", "personal projects", "key projects",
    "publications", "awards", "honors",
    "volunteer", "volunteering", "community involvement",
    "languages", "interests", "references",
]

_EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
_PHONE_RE = re.compile(
    r"(?:\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}"
)
_LINKEDIN_RE = re.compile(
    r"(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+"
)
_GITHUB_RE = re.compile(
    r"(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+"
)


class ParsedResume(BaseModel):
    """Structured data extracted from a parsed resume."""

    model_config = ConfigDict(frozen=True)

    raw_text: str
    file_path: str
    file_format: str
    sections: dict[str, str] = Field(default_factory=dict)
    contact_info: dict[str, str] = Field(default_factory=dict)
    skills: list[str] = Field(default_factory=list)
    word_count: int = 0


class DocumentParser:
    """Async parser for PDF and DOCX resume files.

    Delegates blocking I/O to a thread-pool executor so the event
    loop is never blocked by file reads or library calls.
    """

    _SUPPORTED_FORMATS: frozenset[str] = frozenset({".pdf", ".docx"})

    async def parse(self, file_path: Path) -> ParsedResume:
        """Parse a resume file and extract structured data.

        Args:
            file_path: Path to the PDF or DOCX file.

        Returns:
            ``ParsedResume`` with extracted text and metadata.

        Raises:
            ParseError: If the file cannot be read or parsed.
        """
        path = Path(file_path)
        if not path.exists():
            raise ParseError(str(path), "File does not exist")

        suffix = path.suffix.lower()
        if suffix not in self._SUPPORTED_FORMATS:
            raise ParseError(
                str(path),
                f"Unsupported format '{suffix}'. Supported: {', '.join(self._SUPPORTED_FORMATS)}",
            )

        logger.info("parsing_document", file_path=str(path), format=suffix)
        loop = asyncio.get_running_loop()

        try:
            if suffix == ".pdf":
                raw_text = await loop.run_in_executor(
                    None, self._parse_pdf_sync, path
                )
            else:
                raw_text = await loop.run_in_executor(
                    None, self._parse_docx_sync, path
                )
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError(str(path), str(exc)) from exc

        sections = self._extract_sections(raw_text)
        contact_info = self._extract_contact_info(raw_text)
        skills = self._extract_skills_from_text(raw_text, sections)

        result = ParsedResume(
            raw_text=raw_text,
            file_path=str(path),
            file_format=suffix.lstrip("."),
            sections=sections,
            contact_info=contact_info,
            skills=skills,
            word_count=len(raw_text.split()),
        )
        logger.info(
            "document_parsed",
            file_path=str(path),
            word_count=result.word_count,
            sections_found=len(sections),
            skills_found=len(skills),
        )
        return result

    def _parse_pdf_sync(self, file_path: Path) -> str:
        """Synchronous PDF parsing, preferring pdfplumber over pypdf.

        A raw PDF-text extractor's ``extract_text()`` can return an empty string (not an
        error) on PDFs with certain font-embedding/encoding schemes common to
        Canva/Google-Docs/LaTeX exports — a real CV can produce zero usable text while
        "succeeding". pdfplumber's layout-aware extraction handles those cases; pypdf remains
        the fallback for whatever pdfplumber cannot open (encrypted PDFs, some malformed
        files) so a single library's blind spot no longer determines whether a resume is
        scoreable at all.

        Args:
            file_path: Path to the PDF file.

        Returns:
            Extracted plain text from all pages.
        """
        text = self._parse_pdf_with_pdfplumber(file_path)
        if text:
            return text
        return self._parse_pdf_with_pypdf(file_path)

    def _parse_pdf_with_pdfplumber(self, file_path: Path) -> str:
        """Best-effort pdfplumber extraction. Returns "" on any failure — never raises —
        so the caller can fall back to pypdf without special-casing pdfplumber's errors."""
        try:
            import pdfplumber
        except ImportError:
            return ""

        try:
            with pdfplumber.open(str(file_path)) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
        except Exception as exc:
            logger.debug(
                "pdfplumber_extraction_failed", file_path=str(file_path), error=str(exc)
            )
            return ""
        return "\n".join(p for p in pages if p)

    def _parse_pdf_with_pypdf(self, file_path: Path) -> str:
        """Synchronous PDF parsing with pypdf (fallback extractor).

        Uses ``pypdf`` rather than the unmaintained PyPDF2 (merged into pypdf years ago,
        last released with an unpatched known CVE) — same ``PdfReader``/``extract_text()``
        API, so this is a drop-in swap, not a rewrite.

        Args:
            file_path: Path to the PDF file.

        Returns:
            Extracted plain text from all pages.
        """
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ParseError(
                str(file_path), "pypdf is required for PDF parsing"
            ) from exc

        reader = PdfReader(str(file_path))
        pages: list[str] = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)

        if not pages:
            raise ParseError(str(file_path), "No text content found in PDF")
        return "\n".join(pages)

    def _parse_docx_sync(self, file_path: Path) -> str:
        """Synchronous DOCX parsing with python-docx.

        Reads body paragraphs, table cells, and header/footer paragraphs. Many ATS-styled
        CV templates (two-column layouts, skills grids) lay out content as a Word table
        rather than paragraphs — reading only ``doc.paragraphs`` silently dropped that
        content, which is why table-based resumes could extract to near-empty text.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Extracted plain text from paragraphs, tables, and headers/footers.
        """
        try:
            from docx import Document
        except ImportError as exc:
            raise ParseError(
                str(file_path), "python-docx is required for DOCX parsing"
            ) from exc

        doc = Document(str(file_path))
        chunks: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        chunks.append(cell_text)

        for section in doc.sections:
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    if p.text.strip():
                        chunks.append(p.text.strip())

        if not chunks:
            raise ParseError(str(file_path), "No text content found in DOCX")
        return "\n".join(chunks)

    def _extract_sections(self, text: str) -> dict[str, str]:
        """Extract resume sections by matching known header patterns.

        Args:
            text: Full resume text.

        Returns:
            Mapping of section name to section content.
        """
        lines = text.split("\n")
        sections: dict[str, str] = {}
        current_section: str | None = None
        current_lines: list[str] = []

        for line in lines:
            stripped = line.strip()
            normalized = stripped.lower().rstrip(":")

            if normalized in _SECTION_HEADERS:
                if current_section is not None:
                    sections[current_section] = "\n".join(current_lines).strip()
                current_section = normalized
                current_lines = []
            elif current_section is not None:
                current_lines.append(stripped)

        if current_section is not None:
            sections[current_section] = "\n".join(current_lines).strip()

        return sections

    def _extract_contact_info(self, text: str) -> dict[str, str]:
        """Extract email, phone, LinkedIn, and GitHub URLs from text.

        Args:
            text: Full resume text.

        Returns:
            Mapping of contact type to extracted value.
        """
        info: dict[str, str] = {}

        email_match = _EMAIL_RE.search(text)
        if email_match:
            info["email"] = email_match.group()

        phone_match = _PHONE_RE.search(text)
        if phone_match:
            info["phone"] = phone_match.group()

        linkedin_match = _LINKEDIN_RE.search(text)
        if linkedin_match:
            info["linkedin"] = linkedin_match.group()

        github_match = _GITHUB_RE.search(text)
        if github_match:
            info["github"] = github_match.group()

        return info

    def _extract_skills_from_text(
        self, text: str, sections: dict[str, str]
    ) -> list[str]:
        """Extract skill keywords from resume text.

        Prioritizes the skills section if found; otherwise scans
        the full text for common skill patterns.

        Args:
            text: Full resume text.
            sections: Already-extracted sections dict.

        Returns:
            Deduplicated list of skill strings.
        """
        skills_text = ""
        for key in ("skills", "technical skills", "core competencies", "competencies"):
            if key in sections:
                skills_text = sections[key]
                break

        if not skills_text:
            skills_text = text

        raw_skills: list[str] = []
        for line in skills_text.split("\n"):
            line = line.strip()
            if not line:
                continue
            # Skills often listed with commas, pipes, bullets, or semicolons
            for delimiter in [",", "|", ";", "\u2022", "\u2023", "\u25e6"]:
                if delimiter in line:
                    raw_skills.extend(
                        s.strip() for s in line.split(delimiter) if s.strip()
                    )
                    break
            else:
                # Single skill per line (common in bullet-point resumes)
                cleaned = line.lstrip("-*\u2022\u2023 ").strip()
                if cleaned and len(cleaned) < 60:
                    raw_skills.append(cleaned)

        # Deduplicate while preserving order
        seen: set[str] = set()
        unique: list[str] = []
        for skill in raw_skills:
            lower = skill.lower()
            if lower not in seen and len(skill) > 1:
                seen.add(lower)
                unique.append(skill)

        return unique
