"""Turn extracted CV text into a :class:`ResumeDocument` without discarding anything.

The rule that governs every decision here: **no line is ever dropped.** The parser this
replaces recognised a section only if its heading matched one of 24 exact strings and threw
away everything else, which on a real CV meant throwing away the entire CV. Recognising a
heading is therefore treated as an optimisation, not a gate — a heading we fail to spot costs
us a section boundary, never content.

Heading detection uses how the line is *written* rather than what it is *called*, so a CV is
free to name its sections whatever it likes. Two signals carry it, both chosen for precision
over recall since a missed heading is cheap and a false one fragments a section:

* the line is set in capitals, which is how the large majority of CVs mark a section, or
* the line ends in a colon, which is how most of the rest do.

Lines carrying an email, a URL or a date range are never headings however they are cased —
"Data Scientist → AI Product Manager June 2024 - Present" is a role line, not a section.
"""
# ruff: noqa: RUF001
# Bullet glyphs, en dashes and hyphen bullets appear as data here: these modules exist
# to recognise the characters real CVs are typeset with. ASCII look-alikes would stop
# them matching the documents they are written for.


from __future__ import annotations

import re

from app.core.resume_tailoring.model import (
    BULLET_PREFIX,
    Line,
    LineKind,
    ResumeDocument,
    Section,
    Style,
    make_line,
)

#: Contact-detail patterns. Their presence disqualifies a line from being a heading and marks
#: it as part of the header block.
_EMAIL = re.compile(r"[\w.%+-]+@[\w.-]+\.[A-Za-z]{2,}")
_URL = re.compile(r"(?:https?://|www\.|linkedin\.com|github\.com)", re.I)
_PHONE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")

#: A month or a year range. Role lines carry these; section headings do not.
_DATE = re.compile(
    r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s*\d{2,4}"
    r"|\b(19|20)\d{2}\s*[-–—]\s*((19|20)\d{2}|present|current|now)\b"
    r"|\b(19|20)\d{2}\b",
    re.I,
)

#: Headings are short. Anything longer is a sentence that happens to be capitalised.
_MAX_HEADING_CHARS = 64
_MAX_HEADING_WORDS = 8

#: Section names whose prose is the candidate's positioning statement, and may therefore be
#: re-pointed at a target role. Matched loosely because CVs name it many ways; an unnamed
#: block at the top of the document counts too.
_SUMMARY_HEADING = re.compile(
    r"\b(summary|profile|objective|about\s+me|professional\s+profile|personal\s+statement)\b",
    re.I,
)


def _uppercase_ratio(text: str) -> float:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return 0.0
    return sum(1 for c in letters if c.isupper()) / len(letters)


def is_heading(raw: str) -> bool:
    """Whether a line introduces a new section.

    Deliberately conservative. A false positive splits one section into two and can strand an
    employer line under the wrong heading; a false negative merely merges two sections, and
    every line is still present and still in order.
    """
    text = raw.strip()
    if not text or len(text) > _MAX_HEADING_CHARS:
        return False
    if len(text.split()) > _MAX_HEADING_WORDS:
        return False
    if _EMAIL.search(text) or _URL.search(text) or _PHONE.search(text):
        return False
    if _DATE.search(text):
        return False
    # A heading is a label, not a sentence.
    if text.endswith((".", ",", ";")):
        return False
    if len([c for c in text if c.isalpha()]) < 3:
        return False

    if text.endswith(":"):
        return True
    # "EDUCATION & QUALIFICATIONS" / "WORK & LEADERSHIP EXPERIENCE".
    return _uppercase_ratio(text) >= 0.7


def _is_contact_line(text: str) -> bool:
    return bool(_EMAIL.search(text) or _URL.search(text) or _PHONE.search(text))


def _split_header(lines: list[str]) -> int:
    """How many leading lines are the contact block.

    The name and contact details are immutable and are not part of any section. The block ends
    at the last line within the opening run that still carries a contact detail — after that
    the document has started saying something, and on many CVs (this one included) what it
    says next is an unlabelled summary paragraph that must stay editable.
    """
    last_contact = -1
    for index, raw in enumerate(lines[:8]):
        if _is_contact_line(raw):
            last_contact = index
        elif is_heading(raw):
            break
    # Always keep at least the name line.
    return max(last_contact + 1, 1) if lines else 0


def parse_resume_text(text: str, *, source_format: str = "") -> ResumeDocument:
    """Parse plain CV text, with no typography to preserve."""
    return parse_styled_lines(
        [(ln, Style()) for ln in (text or "").splitlines()], source_format=source_format
    )


def parse_styled_lines(
    pairs: list[tuple[str, Style]], *, source_format: str = "", joined: bool = False
) -> ResumeDocument:
    """Parse CV lines that carry their own typography.

    Every non-blank line of the input appears exactly once in the result, in its original
    order. That property is asserted by the test suite, because it is the one the previous
    implementation violated catastrophically.

    Style travels with the line from here on. Deciding typography at render time instead is
    what made every tailored CV come back in the renderer's own font rather than the
    candidate's — a document that no longer looked like theirs however accurate the words.

    ``joined`` says each pair is already a whole paragraph. The PDF layout reader groups
    wrapped lines itself, from their coordinates, so the text-based re-joining below is skipped
    — it would only be guessing again at what the page already says.
    """
    non_blank = [(raw, st) for raw, st in pairs if raw.strip()]
    if not non_blank:
        return ResumeDocument(source_format=source_format)

    ordered = [raw for raw, _ in non_blank]
    styles = [st for _, st in non_blank]
    header_len = _split_header(ordered)

    doc = ResumeDocument(source_format=source_format)
    counter = 0
    for raw, st in zip(ordered[:header_len], styles[:header_len], strict=True):
        line = make_line(counter, raw)
        # Nothing in the contact block is ever edited, even if it looks like a bullet.
        line.kind = LineKind.FIXED
        line.style = st
        doc.header.append(line)
        counter += 1

    #: The run of content before the first heading. Real CVs open with an unlabelled summary
    #: far more often than not, and the previous parser discarded exactly this.
    current = Section(heading="", order=0)
    order = 0

    for raw, st in zip(ordered[header_len:], styles[header_len:], strict=True):
        if is_heading(raw):
            if current.lines or current.heading:
                doc.sections.append(current)
            order += 1
            current = Section(heading=raw.strip(), order=order, heading_style=st)
            continue
        line = make_line(counter, raw)
        line.style = st
        counter += 1
        current.lines.append(line)

    if current.lines or current.heading:
        doc.sections.append(current)

    if not joined:
        _join_wrapped_lines(doc)
    _protect_skill_tables(doc)
    _classify_prose(doc)
    if not joined:
        _join_prose(doc)
    return doc


#: A heading naming the skills block. Its content is usually a table, flattened by text
#: extraction into rows that superficially look like bullets. Editing those rows would
#: scramble the table, so the whole section is held immutable and keyword additions are
#: routed to experience bullets instead — which is where a recruiter reads them anyway.
#: Words that name a skills block. Substring matching, deliberately: word-boundary
#: patterns kept failing on the plural, and a heading is short enough that a substring
#: test is both sufficient and impossible to mis-escape.
_SKILLS_WORDS = ("skill", "competenc", "technolog", "technical", "highlight",
                 "proficienc", "toolkit", "stack")


def _is_skills_heading(heading: str) -> bool:
    """Whether this heading introduces the skills block."""
    low = heading.lower()
    return any(word in low for word in _SKILLS_WORDS)

#: An entry header: an employer, a role-and-dates line, a project sub-heading. These follow
#: bullets all the time and must never be swallowed into the bullet above them.
_ENTRY_HEADER_HINT = re.compile(
    r"\((?:Technology|Product|Engineering|Data)\s+Team\)|–|—", re.I
)


def _stitch(left: str, right: str) -> str:
    """Join two halves of a wrapped line.

    A trailing hyphen belongs to the word — extraction splits "Azure-based" across lines, and
    both dropping the hyphen ("Azurebased") and padding it ("Azure- based") corrupt a
    technology name that a recruiter and an ATS both search for verbatim.
    """
    left = left.rstrip()
    if left.endswith("-"):
        return f"{left}{right.lstrip()}"
    return f"{left} {right.lstrip()}".strip()


def _looks_like_continuation(previous: Line, raw: str) -> bool:
    """Whether ``raw`` is the wrapped remainder of the line above it.

    PDF text extraction breaks one résumé bullet across several physical lines. Left alone
    each fragment becomes its own line and only the first is editable, so an edit could
    rewrite the opening clause of a bullet and strand the rest — which reads far worse than
    no edit at all.

    Two signals decide it, and they cover the cases a CV actually produces: the previous line
    stopped mid-sentence, or this one starts mid-sentence.
    """
    text = raw.strip()
    if not text or BULLET_PREFIX.match(raw) or is_heading(raw):
        return False
    if _DATE.search(text) or _ENTRY_HEADER_HINT.search(text):
        return False
    previous_open = not previous.text.rstrip().endswith((".", "!", "?", ":", ";"))
    starts_mid_sentence = text[:1].islower()
    return previous_open or starts_mid_sentence


def _join_wrapped_lines(doc: ResumeDocument) -> None:
    """Fold wrapped fragments back into the bullet they belong to."""
    for section in doc.sections:
        merged: list[Line] = []
        for line in section.lines:
            if (
                merged
                and merged[-1].kind is LineKind.BULLET
                and line.kind is LineKind.FIXED
                and _looks_like_continuation(merged[-1], line.text)
            ):
                merged[-1].text = _stitch(merged[-1].text, line.text)
                continue
            merged.append(line)
        section.lines = merged


def _protect_skill_tables(doc: ResumeDocument) -> None:
    """Hold every line of a skills section immutable."""
    for section in doc.sections:
        if section.heading and _is_skills_heading(section.heading):
            for line in section.lines:
                line.kind = LineKind.FIXED


def _join_prose(doc: ResumeDocument) -> None:
    """Fold a wrapped summary paragraph into a single editable line."""
    for section in doc.sections:
        merged: list[Line] = []
        for line in section.lines:
            if merged and merged[-1].kind is LineKind.PROSE and line.kind is LineKind.PROSE:
                merged[-1].text = _stitch(merged[-1].text, line.text)
                continue
            merged.append(line)
        section.lines = merged


def _classify_prose(doc: ResumeDocument) -> None:
    """Mark summary prose editable; leave every other non-bullet line immutable.

    Only a section the CV already has can be edited this way. Nothing here creates a summary,
    and a CV without one simply has no editable prose — which is the correct outcome, not a
    gap to be filled.
    """
    for section in doc.sections:
        is_summary = not section.heading or bool(_SUMMARY_HEADING.search(section.heading))
        if not is_summary:
            continue
        for line in section.lines:
            if line.kind is LineKind.FIXED and not _is_contact_line(line.text):
                line.kind = LineKind.PROSE


#: Fonts are embedded with a subset prefix like "BCDEEE+TimesNewRomanPSMT". Word needs the
#: family name, so the prefix and the PostScript suffixes come off.
_SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
_PS_SUFFIX = re.compile(r"(PSMT|PS-BoldMT|PS-ItalicMT|PS-BoldItalicMT|MT|PS)$")


def _family(fontname: str) -> str:
    """The typeface family a PDF font name refers to."""
    name = _SUBSET_PREFIX.sub("", fontname or "")
    name = name.split(",")[0].split("-")[0]
    name = _PS_SUFFIX.sub("", name)
    # "TimesNewRoman" -> "Times New Roman"; Word matches on the spaced name.
    spaced = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name).strip()
    return spaced or name


def _colour(raw: object) -> tuple[int, int, int] | None:
    """A PDF colour as 0-255 RGB, or None when it is plain black.

    PDFs express colour in several spaces: a single grey value, three RGB components, or four
    CMYK. Treating the one-tuple case as red is the classic mistake here.
    """
    if not isinstance(raw, (list, tuple)) or not raw:
        return None
    values = [float(v) for v in raw]
    if len(values) == 1:
        grey = values[0]
        return None if grey == 0 else (round(grey * 255),) * 3
    if len(values) == 3:
        rgb = tuple(round(v * 255) for v in values)
        return None if rgb == (0, 0, 0) else rgb
    if len(values) == 4:
        c, m, y, k = values
        rgb = tuple(round(255 * (1 - min(1.0, comp + k))) for comp in (c, m, y))
        return None if rgb == (0, 0, 0) else rgb
    return None


def _line_style(line: dict, page_width: float, previous_bottom: float | None) -> Style:
    """Summarise how one extracted line is typeset.

    The dominant value wins for each attribute rather than the first: a heading whose final
    character is punctuation in a different face is still that heading's font.
    """
    from collections import Counter

    chars = line.get("chars") or []
    if not chars:
        return Style()

    fonts = Counter(_family(c.get("fontname", "")) for c in chars if c.get("fontname"))
    sizes = Counter(round(float(c.get("size", 0)), 1) for c in chars)
    colours = Counter(str(c.get("non_stroking_color")) for c in chars)
    bold_chars = sum(1 for c in chars if "bold" in str(c.get("fontname", "")).lower())
    italic_chars = sum(1 for c in chars if "italic" in str(c.get("fontname", "")).lower())

    dominant_colour_key = colours.most_common(1)[0][0]
    swatch = next(
        (c.get("non_stroking_color") for c in chars
         if str(c.get("non_stroking_color")) == dominant_colour_key),
        None,
    )

    x0, x1 = float(line.get("x0", 0)), float(line.get("x1", 0))
    centre_offset = abs(((x0 + x1) / 2) - (page_width / 2))
    align = "center" if centre_offset < page_width * 0.06 and x0 > page_width * 0.2 else "left"

    gap = 0.0
    if previous_bottom is not None:
        gap = max(0.0, float(line.get("top", 0)) - previous_bottom)

    return Style(
        font=fonts.most_common(1)[0][0] if fonts else "",
        size=sizes.most_common(1)[0][0] if sizes else 0.0,
        bold=bold_chars > len(chars) / 2,
        italic=italic_chars > len(chars) / 2,
        colour=_colour(swatch),
        align=align,
        # Only report a gap that is genuinely a paragraph break, not normal leading.
        space_before=round(gap, 1) if gap > 6 else 0.0,
    )


def extract_from_pdf(path: str) -> ResumeDocument:
    """Parse a PDF CV, keeping the typography the candidate actually used."""
    import pdfplumber

    pairs: list[tuple[str, Style]] = []
    lefts: list[float] = []
    rights: list[float] = []
    tops: list[float] = []
    bottoms: list[float] = []
    page_width = page_height = 0.0

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_width, page_height = float(page.width), float(page.height)
            previous_bottom: float | None = None
            for line in page.extract_text_lines():
                text = line.get("text", "")
                if not text.strip():
                    continue
                pairs.append((text, _line_style(line, page_width, previous_bottom)))
                lefts.append(float(line.get("x0", 0)))
                rights.append(float(line.get("x1", 0)))
                tops.append(float(line.get("top", 0)))
                bottoms.append(float(line.get("bottom", 0)))
                previous_bottom = float(line.get("bottom", 0))

    doc = parse_styled_lines(pairs, source_format="pdf")
    if lefts and page_width:
        # The smallest left edge is the text block's margin; the largest right edge gives the
        # mirror. Measured rather than assumed, because Word's one-inch default reflows a CV
        # that was laid out to fit three pages.
        doc.page_margins = (
            round(min(lefts), 1),
            round(max(0.0, page_width - max(rights)), 1),
            round(min(tops), 1),
            round(max(0.0, page_height - max(bottoms)), 1),
        )
    return doc


def extract_from_docx(path: str) -> ResumeDocument:
    """Parse a DOCX CV.

    Table cells are read as well as paragraphs: CVs routinely lay the skills matrix out as a
    table, and reading paragraphs alone loses the entire section.
    """
    from docx import Document as DocxDocument

    document = DocxDocument(path)
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return parse_resume_text("\n".join(parts), source_format="docx")
